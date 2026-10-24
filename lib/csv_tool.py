#!/usr/bin/env python3
"""CSV 数据工具箱 —— 切分 / 去重 / 校验 / 清洗 / 合并 / 转 Excel（自动识别编码与分隔符）"""
import argparse, csv, io, json, os, re, sys
from pathlib import Path

ENCODINGS = ['utf-8-sig', 'utf-8', 'gb18030', 'gbk', 'big5']


def read_text(path: Path):
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='replace'), 'utf-8(replace)'


def sniff(path: Path):
    text, enc = read_text(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        delim = dialect.delimiter
    except Exception:
        delim = ','
    has_header = True
    try:
        has_header = csv.Sniffer().has_header(sample)
    except Exception:
        pass
    return {'编码': enc, '分隔符': 'TAB' if delim == '\t' else delim, '疑似有表头': has_header}


def load_csv(path: Path):
    text, enc = read_text(path)
    info = sniff(path)
    delim = '\t' if info['分隔符'] == 'TAB' else info['分隔符']
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    rows = [r for r in rows if any((c or '').strip() for c in r)]
    return rows, enc, delim


def write_csv(path: Path, rows, enc='utf-8-sig'):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', newline='', encoding=enc) as f:
        w = csv.writer(f)
        for r in rows:
            w.writerow(r)


def safe_out(outdir: Path, name: str, suffix: str):
    p = outdir / (Path(name).stem + suffix)
    i = 1
    while p.exists():
        p = outdir / (Path(name).stem + f'_{i}' + suffix)
        i += 1
    return p


def out_dir(args, src: Path):
    return Path(args.out) if args.out else (src.parent / '_输出')


def iter_files(inp: Path):
    if inp.is_file():
        yield inp
        return
    for p in sorted(inp.rglob('*')):
        if p.is_file() and p.suffix.lower() in ('.csv', '.tsv', '.txt'):
            yield p


def cmd_sniff(args):
    inp = Path(args.input)
    files = list(iter_files(inp))
    if not files:
        return {'ok': False, '错误': '没有找到 CSV/TSV 文件'}
    out = []
    for f in files:
        out.append({'文件': f.name, **sniff(f)})
    return {'ok': True, '文件': out}


def cmd_split(args):
    inp = Path(args.input)
    if inp.is_file():
        files = [inp]
    else:
        files = list(iter_files(inp))
    if not files:
        return {'ok': False, '错误': '没有找到可切分的文件'}
    size = max(1, int(args.rows or 10000))
    outdir = out_dir(args, inp)
    result = []
    for f in files:
        rows, enc, delim = load_csv(f)
        header, body = (rows[0], rows[1:]) if rows else ([], [])
        parts = [body[i:i + size] for i in range(0, len(body), size)] or [[]]
        files_out = []
        for idx, part in enumerate(parts, 1):
            target = outdir / f'{f.stem}-part{idx:02d}.csv'
            write_csv(target, ([header] if header else []) + part)
            files_out.append(str(target))
        result.append({'文件': f.name, '编码': enc, '总行': len(body), '分片': len(parts), '产物': files_out})
    return {'ok': True, '结果': result}


def norm(v):
    return (v or '').strip()


def cmd_dedupe(args):
    inp = Path(args.input)
    if not inp.is_file():
        return {'ok': False, '错误': '去重需要指定单个文件（--input）'}
    rows, enc, delim = load_csv(inp)
    if not rows:
        return {'ok': True, '原行数': 0, '去重后': 0, '删除': 0}
    header, body = rows[0], rows[1:]
    keys = [k.strip() for k in (args.key or '').split(',') if k.strip()]
    idxs = [header.index(k) for k in keys if k in header] if keys else None
    if keys and not idxs:
        return {'ok': False, '错误': f'指定的列不存在：{keys}（表头：{header}）'}
    seen, kept = set(), []
    for r in body:
        k = tuple(norm(r[i]) for i in idxs) if idxs is not None else tuple(norm(c) for c in r)
        if k in seen:
            continue
        seen.add(k)
        kept.append(r)
    outdir = out_dir(args, inp)
    target = outdir / f'{inp.stem}-去重.csv'
    write_csv(target, [header] + kept)
    return {'ok': True, '原行数': len(body), '去重后': len(kept), '删除': len(body) - len(kept), '产物': str(target), '编码': enc}


DATE_RE = re.compile(r'^\d{4}[-/年]\d{1,2}[-/月]\d{1,2}日?$')
EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
PHONE_RE = re.compile(r'^1[3-9]\d{9}$')


def cmd_validate(args):
    inp = Path(args.input)
    files = [inp] if inp.is_file() else list(iter_files(inp))
    if not files:
        return {'ok': False, '错误': '没有找到可校验的文件'}
    required = [c.strip() for c in (args.required or '').split(',') if c.strip()]
    numeric = [c.strip() for c in (args.numeric or '').split(',') if c.strip()]
    dates = [c.strip() for c in (args.date or '').split(',') if c.strip()]
    emails = [c.strip() for c in (args.email or '').split(',') if c.strip()]
    phones = [c.strip() for c in (args.phone or '').split(',') if c.strip()]
    total_errors = 0
    reports = []
    for f in files:
        rows, enc, _ = load_csv(f)
        if not rows:
            continue
        header, body = rows[0], rows[1:]
        col = {name: i for i, name in enumerate(header)}
        issues = []
        for ln, r in enumerate(body, start=2):
            def get(name):
                i = col.get(name)
                return norm(r[i]) if i is not None and i < len(r) else ''
            for c in required:
                if c in col and not get(c):
                    issues.append({'行': ln, '列': c, '问题': '必填为空', '值': ''})
            for c in numeric:
                v = get(c)
                if v and not re.fullmatch(r'-?\d+(\.\d+)?', v):
                    issues.append({'行': ln, '列': c, '问题': '不是数字', '值': v})
            for c in dates:
                v = get(c)
                if v and not DATE_RE.match(v):
                    issues.append({'行': ln, '列': c, '问题': '日期格式不符', '值': v})
            for c in emails:
                v = get(c)
                if v and not EMAIL_RE.match(v):
                    issues.append({'行': ln, '列': c, '问题': '邮箱格式不符', '值': v})
            for c in phones:
                v = get(c)
                if v and not PHONE_RE.match(v):
                    issues.append({'行': ln, '列': c, '问题': '手机号格式不符', '值': v})
        total_errors += len(issues)
        reports.append({'文件': f.name, '行数': len(body), '错误数': len(issues), '错误': issues[:100]})
    outdir = out_dir(args, inp)
    report = outdir / '校验报告.json'
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps({'总错误数': total_errors, '文件': reports}, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'ok': total_errors == 0, '总错误数': total_errors, '文件': reports, '报告': str(report)}


def cmd_clean(args):
    inp = Path(args.input)
    files = [inp] if inp.is_file() else list(iter_files(inp))
    if not files:
        return {'ok': False, '错误': '没有找到可清洗的文件'}
    outdir = out_dir(args, inp)
    result = []
    for f in files:
        rows, enc, _ = load_csv(f)
        if not rows:
            continue
        header = [norm(c) for c in rows[0]]
        body = []
        for r in rows[1:]:
            cells = [norm(c) for c in r]
            if not any(cells):
                continue
            body.append(cells)
        target = outdir / f'{f.stem}-清洗.csv'
        write_csv(target, [header] + body)
        result.append({'文件': f.name, '原行数': len(rows) - 1, '清洗后': len(body), '产物': str(target)})
    return {'ok': True, '结果': result}


def cmd_merge(args):
    inp = Path(args.input)
    if inv := (inp.is_file() or not inp.is_dir()):
        return {'ok': False, '错误': '合并需要 --input 指定目录'}
    files = list(iter_files(inp))
    files = [f for f in files if f.parent == inp or True]
    if len(files) < 2:
        return {'ok': False, '错误': '目录里至少需要 2 个 CSV 文件'}
    merged, order = [], []
    for f in files:
        rows, _, _ = load_csv(f)
        if not rows:
            continue
        header = [norm(c) for c in rows[0]]
        if not order:
            order = header
        idx = {name: i for i, name in enumerate(header)}
        for r in rows[1:]:
            merged.append([norm(r[idx.get(c, -1)]) if c in idx and idx[c] < len(r) else '' for c in order] + [f.name])
    outdir = out_dir(args, inp)
    target = outdir / '合并.csv'
    write_csv(target, [order + (['来源文件'] if args.source_col else [])] + merged)
    return {'ok': True, '文件数': len(files), '行数': len(merged), '列': order, '产物': str(target)}


def cmd_convert(args):
    inp = Path(args.input)
    files = [inp] if inp.is_file() else list(iter_files(inp))
    if not files:
        return {'ok': False, '错误': '没有找到可转换的文件'}
    fmt = (args.format or 'xlsx').lower()
    if fmt not in ('xlsx', 'csv'):
        return {'ok': False, '错误': '--format 只支持 xlsx / csv'}
    outdir = out_dir(args, inp)
    out = []
    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
        except ImportError:
            return {'ok': False, '错误': '需要 openpyxl：python3 -m pip install openpyxl', '缺依赖': 'openpyxl'}
        for f in files:
            rows, _, _ = load_csv(f)
            wb = Workbook()
            ws = wb.active
            for r in rows:
                ws.append(r)
            outdir.mkdir(parents=True, exist_ok=True)
            target = outdir / f'{f.stem}.xlsx'
            i = 1
            while target.exists():
                target = outdir / f'{f.stem}_{i}.xlsx'
                i += 1
            wb.save(target)
            out.append({'文件': f.name, '产物': str(target), '行数': len(rows)})
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {'ok': False, '错误': '需要 openpyxl：python3 -m pip install openpyxl', '缺依赖': 'openpyxl'}
        for f in files:
            if f.suffix.lower() not in ('.xlsx', '.xlsm'):
                continue
            wb = load_workbook(f, data_only=True, read_only=True)
            for ws in wb.worksheets:
                rows = [[('' if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
                target = outdir / f'{f.stem}-{ws.title}.csv'
                write_csv(target, rows)
                out.append({'文件': f.name, '工作表': ws.title, '产物': str(target), '行数': len(rows)})
    return {'ok': True, '格式': fmt, '结果': out}


def cmd_demo(args):
    import tempfile
    d = Path(tempfile.mkdtemp(prefix='dsh-csv-lab-demo-'))
    src = d / '输入'
    src.mkdir()
    rows = [['姓名', '邮箱', '手机号', '金额', '日期'],
            ['张三', 'a@example.com', '13800138000', '100', '2026-01-01'],
            ['李四', 'bad-email', '12345', 'abc', '2026/13/40'],
            ['张三', 'a@example.com', '13800138000', '100', '2026-01-01'],
            ['王五', '  w@example.com  ', '13900139000', '200.5', '2026-02-02']]
    p1 = src / '样本1.csv'
    write_csv(p1, rows)
    p2 = src / '样本2.csv'
    write_csv(p2, [['姓名', '邮箱', '手机号', '金额', '日期'], ['赵六', 'z@example.com', '13700137000', '300', '2026-03-03']])
    # GBK 编码样本
    p3 = src / 'GBK样本.csv'
    p3.write_bytes('编号,名称\n1,测试甲\n2,测试乙\n'.encode('gbk'))
    out = d / '输出'
    steps = []
    steps.append(('识别编码分隔符', cmd_sniff(argparse.Namespace(input=str(src)))))
    steps.append(('切分', cmd_split(argparse.Namespace(input=str(p1), out=str(out / 'split'), rows='2'))))
    steps.append(('去重', cmd_dedupe(argparse.Namespace(input=str(p1), out=str(out / 'dedupe'), key=''))))
    steps.append(('校验', cmd_validate(argparse.Namespace(input=str(src), out=str(out / 'validate'), required='姓名,邮箱', numeric='金额', date='日期', email='邮箱', phone='手机号'))))
    steps.append(('清洗', cmd_clean(argparse.Namespace(input=str(src), out=str(out / 'clean')))))
    steps.append(('合并', cmd_merge(argparse.Namespace(input=str(src), out=str(out / 'merge'), source_col=True))))
    steps.append(('转 Excel', cmd_convert(argparse.Namespace(input=str(src), out=str(out / 'xlsx'), format='xlsx'))))
    lines = [f'CSV 数据工具箱 · 演示（--demo）', f'演示目录 {d}', '']
    for name, res in steps:
        ok = res.get('ok', True)
        lines.append(('✓' if ok else '✗') + ' ' + name)
    lines.append('')
    lines.append('产物：')
    for p in sorted(out.rglob('*')):
        if p.is_file():
            lines.append('  ' + str(p))
    return '\n'.join(lines)


def main():
    ap = argparse.ArgumentParser(add_help=False)
    ap.add_argument('action', nargs='?')
    ap.add_argument('--input', '-i'); ap.add_argument('--out', '-o')
    ap.add_argument('--rows'); ap.add_argument('--key'); ap.add_argument('--required')
    ap.add_argument('--numeric'); ap.add_argument('--date'); ap.add_argument('--email'); ap.add_argument('--phone')
    ap.add_argument('--format'); ap.add_argument('--source-col', action='store_true')
    ap.add_argument('--json', action='store_true'); ap.add_argument('--demo', action='store_true')
    ap.add_argument('--help', '-h', action='store_true')
    args = ap.parse_args()
    if args.help or (not args.action and not args.demo):
        print('CSV 数据工具箱\n\n用法：\n  dsh-csv-lab <action> --input <文件|目录> [--out <目录>] [--json]\n  dsh-csv-lab --demo\n  dsh-csv-lab --help\n\n子命令：\n  sniff     识别编码/分隔符/是否有表头\n  split     按行数切分（--rows，默认 10000）\n  dedupe    去重（--key 指定列，默认整行）\n  validate  字段校验（--required/--numeric/--date/--email/--phone）\n  clean     去空格、删空行、规范表头\n  merge     合并目录内多个 CSV（按表头映射，--source-col 加来源列）\n  convert   转 Excel（--format xlsx）或 Excel 转 CSV（--format csv）')
        return 0
    if args.demo:
        print(cmd_demo(args)); return 0
    if not args.input:
        print('缺 --input（或 --demo）', file=sys.stderr); return 2
    inp = Path(args.input)
    if not inp.exists():
        print(f'输入不存在：{inp}', file=sys.stderr); return 2
    ops = {'sniff': cmd_sniff, 'split': cmd_split, 'dedupe': cmd_dedupe, 'validate': cmd_validate, 'clean': cmd_clean, 'merge': cmd_merge, 'convert': cmd_convert}
    fn = ops.get(args.action)
    if fn is None:
        print(f'未知子命令：{args.action}（可用：{", ".join(ops)}）', file=sys.stderr); return 2
    try:
        r = fn(args)
    except Exception as e:
        print(f'执行失败：{e}', file=sys.stderr); return 3
    if r.get('缺依赖'):
        print(f"缺依赖：{r['缺依赖']}。安装：python3 -m pip install {r['缺依赖']}", file=sys.stderr)
        return 3
    if args.json:
        print(json.dumps(r, ensure_ascii=False, indent=2))
    else:
        for k, v in r.items():
            if k in ('文件', '结果', '错误') and isinstance(v, list):
                for item in v[:10]:
                    print(json.dumps(item, ensure_ascii=False))
                if len(v) > 10:
                    print(f'… 共 {len(v)} 项')
            else:
                print(f'{k}: {v}')
    return 0 if r.get('ok', True) else 1


if __name__ == '__main__':
    sys.exit(main())
